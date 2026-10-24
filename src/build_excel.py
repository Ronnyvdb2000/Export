"""
build_excel.py – Ruwe Excel → gestructureerd werkboek met tabbladen
"""

import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TITLE_LINE1 = "VLIF CODELIJST"
TITLE_LINE2 = ""
TITLE_LINE3 = ""

# Mapping van sectienummer naar tabbladnaam en volledige titel
TAB_MAP = {
    "2.1": ("Rundveestal", "2.1 Rundveestal"),
    "2.2": ("Varkensstal", "2.2 Varkensstal"),
    "2.3": ("Pluimveestal", "2.3 Pluimveestal"),
    "2.4": ("Vleeskalverstal", "2.4 Vleeskalverstal"),
    "2.5": ("Paardenstal", "2.5 Paardenstal"),
    "2.6": ("Schapenstal", "2.6 Schapenstal"),
    "2.7": ("Geitenstal", "2.7 Geitenstal"),
    "2.8": ("Hertenstal", "2.8 Hertenstal"),
    "2.9": ("Konijnenstal", "2.9 Konijnenstal"),
    "2.10": ("Serre", "2.10 Serre"),
    "2.11": ("Loods", "2.11 Loods"),
    "2.12": ("Onroerend Buiten", "2.12 Onroerend Buiten Stal/ Loods/ Serre"),
    "2.13": ("Machines en Varia", "2.13 Machines en Varia"),
    "3.1": ("AERi RV", "3.1 AERi Rundveestal"),
    "3.2": ("AERi V", "3.2 AERi Varkensstal"),
    "3.3": ("AERi Pl", "3.3 AERi Pluimveestal"),
    "3.4": ("AERi Vleeskalver", "3.4 AERi Vleeskalverstal"),
    "3.5": ("AERi OG", "3.5 AERi Onroerend Buiten Nieuwbouwstal"),
    "3.6": ("AERi Roerend", "3.6 Roerende AEA-technieken"),
    "4.1": ("NPI Onroerend", "4.1 Onroerend Buiten Stal/ Loods/ Serre"),
    "4.2": ("NPI Machines", "4.2 NPI Machines en Varia"),
}

HEADERS = ["Subfilter", "Code (%)", "Omschrijving", "Extra info",
           "Kenmerk", "Eenheid", "Offerte"]

C_DARK   = "1F4E79"
C_MID    = "2E75B6"
C_LIGHT  = "4472C4"
C_SEC    = "D6E4F0"
C_ALT    = "EBF0F7"
C_WHITE  = "FFFFFF"
C_TOC    = "2E75B6"
C_TOC_LT = "DEEAF1"


def _brd():
    t = Side(style="thin", color="B8CCE4")
    return Border(top=t, bottom=t, left=t, right=t)


def _fix(df):
    df = df.copy().fillna("")
    for col in ("eenheid", "kenmerk"):
        if col in df.columns:
            mask = df[col].astype(str).str.upper().str.contains("OFFERTE", na=False)
            df.loc[mask, "offerte"] = "OFFERTE"
            df.loc[mask, col] = df.loc[mask, col].astype(str).str.replace(
                r"\s*/?\s*OFFERTE\s*/?", "", regex=True, flags=re.IGNORECASE
            ).str.strip(" /")
    return df


def _build_tabs(df):
    """Detecteer tabbladen op basis van het sectienummer (bv. "2.1")."""
    PREFIX_RE = re.compile(r"^(\d+\.\d+)")
    seen = {}
    order = []
    
    for hoofd in df["hoofdstuk"]:
        if hoofd:
            m = PREFIX_RE.match(str(hoofd))
            if m:
                sectie = m.group(1)
                if sectie not in seen:
                    seen[sectie] = True
                    order.append(sectie)
                    
    tabs = []
    for sectie in order:
        if sectie in TAB_MAP:
            tab_name, full_title = TAB_MAP[sectie]
        else:
            tab_name = f"Sectie {sectie}"
            full_title = f"Sectie {sectie}"
        tabs.append((tab_name, sectie, full_title))
    return tabs


def _title_rows(ws, tab_title, ncols):
    ws.append([TITLE_LINE1])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1)
    c.font = Font(bold=True, name="Arial", size=15, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=C_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 22

    ws.append([TITLE_LINE2])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(2, 1)
    c.font = Font(name="Arial", size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=C_MID)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 16

    ws.append([f"{TITLE_LINE3}   |   {tab_title}"])
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    c = ws.cell(3, 1)
    c.font = Font(bold=True, name="Arial", size=9, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=C_LIGHT)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 14

    ws.append([])
    ws.row_dimensions[4].height = 5


def _data_sheet(ws, tab_df, tab_title):
    nc = len(HEADERS)
    _title_rows(ws, tab_title, nc)

    hdr_fill = PatternFill("solid", fgColor=C_DARK)
    hdr_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    dat_font = Font(name="Arial", size=9)
    brd = _brd()
    wrap_cols = {3, 4, 5}

    ws.append(HEADERS)
    for c in range(1, nc + 1):
        cell = ws.cell(5, c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.border = brd
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[5].height = 24

    for i, w in enumerate([30, 7, 42, 52, 28, 12, 9], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row_idx  = 6
    alt      = False
    prev_sf  = None

    for _, rec in tab_df.iterrows():
        sf = rec.get("subfilter", "")

        if sf != prev_sf and sf:
            prev_sf = sf
            alt = False

        fill = PatternFill("solid", fgColor=C_ALT if alt else C_WHITE)
        values = [
            sf,
            rec.get("code (%)", rec.get("code", "")),
            rec.get("omschrijving", ""),
            rec.get("extra info", ""),
            rec.get("kenmerk", ""),
            rec.get("eenheid", ""),
            rec.get("offerte", ""),
        ]
        for ci, val in enumerate(values, start=1):
            cell = ws.cell(row_idx, ci, val or "")
            cell.font      = dat_font
            cell.fill      = fill
            cell.border    = brd
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(ci in wrap_cols))
        ws.row_dimensions[row_idx].height = 15
        row_idx += 1
        alt = not alt

    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False


def _toc(ws, tabs_present):
    ws.title = "Inhoudsopgave"
    ws.sheet_view.showGridLines = False
    nc = 3
    _title_rows(ws, "Inhoudsopgave", nc)

    ws.append(["#", "Tabblad", "Sectie"])
    for c in range(1, 4):
        cell = ws.cell(5, c)
        cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        cell.fill = PatternFill("solid", fgColor=C_TOC)
        cell.border = _brd()
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 22

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 48

    lnk_font = Font(name="Arial", size=10, color=C_DARK,
                    underline="single", bold=True)
    txt_font  = Font(name="Arial", size=10)
    brd = _brd()

    for i, (tab_name, _, full_title) in enumerate(tabs_present, start=1):
        row = 5 + i
        bg  = C_TOC_LT if i % 2 == 0 else C_WHITE
        fill = PatternFill("solid", fgColor=bg)

        n = ws.cell(row, 1, i)
        n.font = txt_font; n.fill = fill; n.border = brd
        n.alignment = Alignment(horizontal="center", vertical="center")

        lnk = ws.cell(row, 2, tab_name)
        lnk.font = lnk_font; lnk.fill = fill; lnk.border = brd
        lnk.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        from openpyxl.worksheet.hyperlink import Hyperlink as XlHyperlink
        safe_name = tab_name.replace("'", "''")
        lnk.hyperlink = XlHyperlink(
            ref=lnk.coordinate,
            location=f"'{safe_name}'!A1",
            tooltip=full_title
        )

        t = ws.cell(row, 3, full_title)
        t.font = txt_font; t.fill = fill; t.border = brd
        t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row].height = 18

    ws.freeze_panes = "A6"


def build(input_xlsx, output_xlsx):
    global TITLE_LINE1, TITLE_LINE2, TITLE_LINE3
    print(f"📖 Inlezen: {input_xlsx}")
    df = pd.read_excel(input_xlsx, dtype=str)
    df = _fix(df)
    
    try:
        meta = pd.read_excel(input_xlsx, sheet_name="_meta", dtype=str)
        titles = meta["titel"].fillna("").tolist()
        if len(titles) >= 1 and titles[0]: TITLE_LINE1 = titles[0]
        if len(titles) >= 2 and titles[1]: TITLE_LINE2 = titles[1]
        if len(titles) >= 3 and titles[2]: TITLE_LINE3 = titles[2]
        print(f"   Titel: {TITLE_LINE1} | {TITLE_LINE2}")
    except Exception:
        pass

    wb = Workbook()
    wb.remove(wb.active)
    toc_ws = wb.create_sheet("Inhoudsopgave", 0)

    tabs = _build_tabs(df)
    tabs_present = []
    for tab_name, sectie, full_title in tabs:
        # Filter: match op "X.Y." of "X.Y " (dus 2.4 matcht met 2.4.1 en met 2.4 Vleeskalverstal)
        tab_df = df[df["hoofdstuk"].str.match(rf"^{sectie}(\s|\.)", na=False)].copy()
        if tab_df.empty:
            print(f"   ⚠️  Leeg: {sectie}")
            continue
        ws = wb.create_sheet(tab_name)
        _data_sheet(ws, tab_df, full_title)
        tabs_present.append((tab_name, sectie, full_title))
        print(f"   ✅ {tab_name:22s}  {len(tab_df):4d} rijen")

    _toc(toc_ws, tabs_present)
    wb.save(output_xlsx)
    print(f"\n🎉 Opgeslagen: {output_xlsx}")


if __name__ == "__main__":
    import argparse
    p = argparse.ArgumentParser()
    p.add_argument("--input",  required=True)
    p.add_argument("--output", required=True)
    a = p.parse_args()
    build(a.input, a.output)
